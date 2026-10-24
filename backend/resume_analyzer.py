"""
Step 9: Resume analyzer (rule-based, no database).

1. extract_text()    - reads text from a PDF or DOCX file
2. detect_sections() - finds Skills / Education / Projects / Experience / Certifications
3. analyze()         - skills found, missing skills for the target career, weak sections,
                       keyword suggestions, and a project-defined resume score
4. suggest_projects()- projects that would add the missing skills to the resume

The score is a project-defined indicator. It is NOT a real ATS score or a hiring prediction.
"""
import io
import re


# ----------------------------------------------------------------- FILE READING
def extract_text(filename, data):
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf"):
            try:
                import pdfplumber
                with pdfplumber.open(io.BytesIO(data)) as pdf:
                    text = "\n".join((page.extract_text() or "") for page in pdf.pages)
            except ImportError:
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(data))
                text = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif name.endswith(".docx"):
            import docx
            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for table in d.tables:
                for row in table.rows:
                    for cell in row.cells:
                        parts.append(cell.text)
            text = "\n".join(parts)
        else:
            raise ValueError("Only PDF and DOCX files are supported")
    except ValueError:
        raise
    except ImportError:
        raise ValueError("Server is missing a library. Run: pip install pdfplumber python-docx")
    except Exception:
        raise ValueError("Could not open this file. Please upload a valid PDF or DOCX resume")

    text = text.strip()
    if len(text) < 30:
        raise ValueError("Could not read any text. Scanned or image-only resumes are not supported; "
                         "please upload a text-based PDF or a DOCX file")
    return text


# -------------------------------------------------------------------- SECTIONS
CORE_SECTIONS = ["Skills", "Education", "Projects", "Experience", "Certifications"]

_HEADINGS = {
    "Skills": ["skills", "technical skills", "key skills", "core skills", "core competencies",
               "skill set", "technical proficiency", "technologies", "tools and technologies"],
    "Education": ["education", "academic background", "academics", "qualification", "qualifications",
                  "educational qualification", "educational qualifications", "academic qualification"],
    "Projects": ["projects", "academic projects", "personal projects", "project work",
                 "key projects", "major projects", "mini projects"],
    "Experience": ["experience", "work experience", "professional experience", "internship",
                   "internships", "internship experience", "employment", "work history", "training"],
    "Certifications": ["certifications", "certification", "certificates", "courses",
                       "certifications and courses", "courses and certifications",
                       "licenses and certifications", "online courses"],
    "Other": ["summary", "professional summary", "objective", "career objective", "profile",
              "about me", "declaration", "hobbies", "interests", "languages", "achievements",
              "awards", "references", "personal details", "contact", "extracurricular activities"],
}
HEADING_LOOKUP = {h: sec for sec, names in _HEADINGS.items() for h in names}


def detect_sections(text):
    sections, current = {}, None
    for line in text.splitlines():
        s = line.strip()
        if not s:
            continue
        norm = re.sub(r"[^a-z& ]", " ", s.lower()).replace("&", "and")
        norm = re.sub(r"\s+", " ", norm).strip()
        if len(s) <= 40 and norm in HEADING_LOOKUP:
            current = HEADING_LOOKUP[norm]
            sections.setdefault(current, [])
            continue
        if current:
            sections[current].append(s)
    return {k: "\n".join(v) for k, v in sections.items()}


# ---------------------------------------------------------------------- SKILLS
SKILL_PATTERNS = {
    "Python": [r"\bpython\b"],
    "SQL": [r"\bsql\b", r"mysql", r"postgres", r"sqlite", r"pl/sql"],
    "Excel": [r"\bexcel\b", r"google sheets", r"spreadsheet", r"pivot"],
    "Statistics": [r"statistic", r"probability", r"hypothesis", r"regression analysis"],
    "Data Visualization": [r"visuali[sz]ation", r"matplotlib", r"seaborn", r"plotly", r"dashboard"],
    "Power BI": [r"power ?bi", r"tableau", r"\bdax\b"],
    "Machine Learning": [r"machine learning", r"\bml\b", r"scikit", r"sklearn", r"random forest"],
    "Deep Learning": [r"deep learning", r"neural network", r"tensorflow", r"pytorch", r"keras", r"\bcnn\b"],
    "DSA": [r"\bdsa\b", r"data structures", r"algorithms", r"leetcode", r"competitive programming"],
    "JavaScript": [r"javascript", r"\bjs\b", r"node\.?js", r"typescript", r"\bes6\b"],
    "HTML/CSS": [r"\bhtml5?\b", r"\bcss3?\b", r"bootstrap", r"tailwind"],
    "React": [r"\breact(\.?js)?\b"],
    "Backend/APIs": [r"rest ?apis?", r"\bapis?\b", r"fastapi", r"flask", r"django", r"spring boot", r"backend"],
    "Git/GitHub": [r"\bgit\b", r"github", r"gitlab"],
    "Linux": [r"linux", r"ubuntu", r"\bbash\b", r"shell script"],
    "Networking": [r"networking", r"tcp/ip", r"\bdns\b", r"\bccna\b", r"wireshark"],
    "Cloud Computing": [r"\baws\b", r"\bazure\b", r"\bgcp\b", r"google cloud", r"cloud computing"],
    "Docker": [r"docker", r"kubernetes", r"container"],
    "CI/CD": [r"ci/cd", r"jenkins", r"github actions"],
    "Cyber Security": [r"cyber", r"security", r"penetration", r"ethical hacking", r"owasp", r"\bnmap\b"],
    "UI/UX Design": [r"figma", r"ui/ux", r"\bux\b", r"wireframe", r"prototyp", r"adobe xd"],
    "Communication": [r"communicat", r"presentation", r"public speaking"],
    "Business Analysis": [r"business analy", r"requirement", r"stakeholder", r"\bbrd\b", r"user stor"],
}
_COMPILED = {s: [re.compile(p, re.I) for p in ps] for s, ps in SKILL_PATTERNS.items()}

SUGGESTED_KEYWORDS = {
    "Python": ["Python", "Pandas", "NumPy"],
    "SQL": ["SQL", "joins", "PostgreSQL / MySQL"],
    "Excel": ["Excel", "pivot tables", "VLOOKUP"],
    "Statistics": ["statistics", "hypothesis testing", "regression"],
    "Data Visualization": ["data visualization", "Matplotlib", "dashboards"],
    "Power BI": ["Power BI", "Tableau", "DAX"],
    "Machine Learning": ["machine learning", "scikit-learn", "model evaluation"],
    "Deep Learning": ["deep learning", "TensorFlow / PyTorch", "neural networks"],
    "DSA": ["data structures", "algorithms", "problem solving"],
    "JavaScript": ["JavaScript", "ES6", "DOM"],
    "HTML/CSS": ["HTML5", "CSS3", "responsive design"],
    "React": ["React", "hooks", "components"],
    "Backend/APIs": ["REST APIs", "FastAPI / Flask", "authentication"],
    "Git/GitHub": ["Git", "GitHub", "version control"],
    "Linux": ["Linux", "Bash", "command line"],
    "Networking": ["TCP/IP", "DNS", "HTTP"],
    "Cloud Computing": ["AWS / Azure / GCP", "cloud deployment"],
    "Docker": ["Docker", "containers"],
    "CI/CD": ["CI/CD", "GitHub Actions", "Jenkins"],
    "Cyber Security": ["OWASP Top 10", "Nmap", "vulnerability assessment"],
    "UI/UX Design": ["Figma", "wireframes", "prototyping"],
    "Communication": ["presentations", "documentation", "team communication"],
    "Business Analysis": ["requirements gathering", "user stories", "stakeholder management"],
}


def find_skills(text, sections):
    """returns {skill: [sections where it appears]} ('Other' = outside a detected section)."""
    found = {}
    for skill, patterns in _COMPILED.items():
        if not any(p.search(text) for p in patterns):
            continue
        where = [sec for sec in CORE_SECTIONS
                 if sec in sections and any(p.search(sections[sec]) for p in patterns)]
        found[skill] = where or ["Other"]
    return found


# -------------------------------------------------------------------- ANALYSIS
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE = re.compile(r"(\+?\d[\d\s\-()]{8,}\d)")
_MIN_WORDS = {"Education": 6, "Projects": 25, "Experience": 20, "Certifications": 3}


def _priority(imp):
    return "High" if imp >= 4 else ("Medium" if imp == 3 else "Low")


def analyze(text, career_name, weights):
    sections = detect_sections(text)
    found = find_skills(text, sections)
    words = len(text.split())

    # sections
    section_info, weak, missing_sections = {}, [], []
    for sec in CORE_SECTIONS:
        content = sections.get(sec)
        if content is None:
            section_info[sec] = {"status": "Missing", "words": 0}
            missing_sections.append(sec)
            continue
        n = len(content.split())
        if sec == "Skills":
            n_sk = sum(1 for s, w in found.items() if sec in w)
            status = "Good" if n_sk >= 3 else "Weak"
        else:
            status = "Good" if n >= _MIN_WORDS[sec] else "Weak"
        section_info[sec] = {"status": status, "words": n}
        if status == "Weak":
            weak.append(sec)

    contact = {
        "email": bool(_EMAIL.search(text)),
        "phone": bool(_PHONE.search(text)),
        "linkedin": "linkedin" in text.lower(),
        "github": "github" in text.lower(),
    }

    # career match
    total_imp = sum(weights.values())
    matched = [s for s in weights if s in found]
    match_pct = sum(weights[s] for s in matched) / total_imp * 100 if total_imp else 0
    missing = sorted(((s, i) for s, i in weights.items() if s not in found), key=lambda x: -x[1])
    missing_skills = [{
        "skill": s, "importance": i, "priority": _priority(i),
        "suggested_keywords": SUGGESTED_KEYWORDS.get(s, [s]),
    } for s, i in missing]

    # score (project-defined indicator, max 100)
    sec_pts = sum(8 if section_info[s]["status"] == "Good" else 4 if section_info[s]["status"] == "Weak" else 0
                  for s in CORE_SECTIONS)
    contact_pts = (8 if contact["email"] else 0) + (6 if contact["phone"] else 0) \
        + (6 if (contact["linkedin"] or contact["github"]) else 0)
    score = round(sec_pts + match_pct * 0.4 + contact_pts)

    # suggestions
    tips = []
    absent = [k for k in ("email", "phone") if not contact[k]]
    if absent:
        tips.append("Add your " + " and ".join(absent) + " in the header of the resume.")
    if not (contact["linkedin"] or contact["github"]):
        tips.append("Add a LinkedIn or GitHub link so recruiters can see your work.")
    for sec in missing_sections:
        if sec == "Certifications":
            tips.append("Add a Certifications / Courses section if you have completed any online courses (optional).")
        else:
            tips.append(f"Add a section for {sec}: it is missing or its heading was not recognised.")
    for sec in weak:
        if sec == "Projects":
            tips.append("Projects section is weak. Describe each project in 2-3 lines: what you built, "
                        "the tools you used and the result.")
        elif sec == "Skills":
            tips.append("Skills section is weak. List your tools and technologies clearly in one place.")
        elif sec == "Experience":
            tips.append("Experience section is short. Add internships, training or club work with what you did and achieved.")
        else:
            tips.append(f"{sec} section is very short. Add more detail.")
    high_missing = [m["skill"] for m in missing_skills if m["priority"] == "High"]
    if high_missing:
        tips.append(f"Your resume does not mention {', '.join(high_missing)}, which matter for {career_name}. "
                    "Add them only if you have really worked with them; otherwise learn them first "
                    "(see your skill gap and roadmap).")
    list_only = [s for s in matched if found[s] == ["Skills"]]
    if list_only and section_info["Projects"]["status"] != "Missing":
        tips.append(f"{', '.join(list_only[:4])} appear only in your skills list. Show them in action by "
                    "mentioning them in a project or experience.")
    if words < 150:
        tips.append("The resume is very short. Add more detail about projects and skills.")
    elif words > 900:
        tips.append("The resume is long. Try to keep it to 1-2 pages.")

    return {
        "career_name": career_name,
        "resume_score": score,
        "score_breakdown": {
            "sections_out_of_40": sec_pts,
            "career_keywords_out_of_40": round(match_pct * 0.4),
            "contact_out_of_20": contact_pts,
        },
        "word_count": words,
        "contact": contact,
        "sections": section_info,
        "weak_sections": weak,
        "missing_sections": missing_sections,
        "skills_found": [{"skill": s, "where": w} for s, w in sorted(found.items())],
        "career_keyword_match_percent": round(match_pct, 1),
        "matched_career_skills": matched,
        "missing_skills": missing_skills,
        "suggestions": tips,
        "note": "The resume score is a project-defined indicator based on sections, career keywords and "
                "contact details. It is not a real ATS score or a hiring prediction.",
    }


def suggest_projects(projects, missing_skill_names, all_skill_names, top_n=3):
    """Projects that would add the most missing skills to the resume."""
    diff = {"Beginner": 0, "Intermediate": 1, "Advanced": 2}
    out = []
    for p in projects:
        tokens = [t.strip() for t in (p["required_skills"] or "").split(",") if t.strip()]
        tokens = [{"ML": "Machine Learning", "APIs": "Backend/APIs"}.get(t, t) for t in tokens]
        adds = [t for t in tokens if t in all_skill_names and t in missing_skill_names]
        out.append({**p, "adds_to_resume": adds})
    out.sort(key=lambda p: (-len(p["adds_to_resume"]), diff.get(p["difficulty"], 9)))
    return out[:top_n]